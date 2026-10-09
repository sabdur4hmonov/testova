"""AUDIT F08: DOCX archive limits (app/services/docx_limits.py).

Step 0 on the unchanged code: a 77 KB DOCX whose document.xml carried 40 MiB of
whitespace made docx_has_renderable_shapes allocate ~89 MB (on the event loop)
and process_file handed it to the renderer instead of rejecting it.

These tests prove: real Word files (text, picture, hyperlink, attached
template) still pass; every unsafe archive shape is rejected before any reader
expands it; and the pipeline turns a rejection into "file_rejected".
"""
import io
import stat
import struct
import tracemalloc
import zipfile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

from app.services import docx_limits, file_processor, pipeline
from app.services.docx_limits import InvalidDocx, validate_docx
from app.services.file_processor import PageImage


def _real_docx(*, picture=False, hyperlink=False, template=False) -> bytes:
    doc = Document()
    doc.add_paragraph("1. 2 + 2 = ?  A) 3  B) 4  C) 5  D) 6")
    if picture:
        img = io.BytesIO()
        Image.new("RGB", (40, 30), "red").save(img, format="PNG")
        img.seek(0)
        doc.add_picture(img)
    if hyperlink:
        doc.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    if template:
        doc.part._settings_part.relate_to(
            "file:///C:/Templates/Normal.dotm",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
            is_external=True,
        )
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _rebuild(blob: bytes, *, replace=None, extra=()) -> bytes:
    """Copy a DOCX, optionally replacing part bytes and appending raw entries."""
    replace = replace or {}
    src = zipfile.ZipFile(io.BytesIO(blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename in replace:
                data = replace[info.filename](data)
            z.writestr(info.filename, data)
        for item in extra:
            if isinstance(item, zipfile.ZipInfo):
                z.writestr(item, b"x")
            else:
                z.writestr(*item)
    return out.getvalue()


def _bomb(pad: int = 40 * 1024 * 1024) -> bytes:
    return _rebuild(_real_docx(), replace={"word/document.xml": lambda d: d + b"\n" * pad})


# ── Real documents keep working ──────────────────────────────────────────────

@pytest.mark.parametrize("kw", [
    {}, {"picture": True}, {"hyperlink": True}, {"template": True},
])
def test_real_word_documents_pass(kw):
    blob = _real_docx(**kw)
    validate_docx(blob)
    pages, text = file_processor.docx_to_images(blob)
    assert pages and "2 + 2" in text


def test_real_picture_is_still_extracted():
    assert len(file_processor.docx_extract_ordered_images(_real_docx(picture=True))) == 1


# ── The Step 0 bomb ──────────────────────────────────────────────────────────

def test_whitespace_bomb_is_rejected_without_expanding_it():
    blob = _bomb()
    assert len(blob) < 200 * 1024                      # tiny upload...
    with pytest.raises(InvalidDocx):
        validate_docx(blob)
    tracemalloc.start()
    assert file_processor.docx_has_renderable_shapes(blob) is False
    _cur, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    assert peak < 4 * 1024 * 1024                      # ...never expanded whole
    with pytest.raises(InvalidDocx):
        file_processor.docx_to_images(blob)
    with pytest.raises(InvalidDocx):
        file_processor.docx_extract_text(blob)
    assert file_processor.docx_extract_ordered_images(blob) == []


def test_ratio_bomb_under_the_size_limits_is_rejected():
    # 5 MiB of newlines: under every size limit, but ~1000x compression.
    with pytest.raises(InvalidDocx, match="ratio"):
        validate_docx(_bomb(5 * 1024 * 1024))


def test_libreoffice_never_sees_a_rejected_archive(monkeypatch):
    def _spawn(*a, **k):
        raise AssertionError("soffice must not run")
    monkeypatch.setattr(file_processor.subprocess, "run", _spawn)
    assert file_processor.docx_to_pdf(_bomb()) is None


async def test_pipeline_rejects_before_any_reader(monkeypatch):
    statuses = []

    async def _status(pid, status, error=None):
        statuses.append((status, error))

    def _never(*a, **k):
        raise AssertionError("reader must not run on a rejected DOCX")

    monkeypatch.setattr(pipeline, "prune_debug_crops", lambda *a, **k: 0)
    monkeypatch.setattr(pipeline, "_set_project_status", _status)
    monkeypatch.setattr(pipeline, "docx_has_renderable_shapes", _never)
    monkeypatch.setattr(pipeline, "docx_to_pdf", _never)
    monkeypatch.setattr(pipeline, "docx_to_images", _never)

    res = await pipeline.process_file(_bomb(), "docx", "pid")

    assert res.status == "file_rejected"
    assert statuses and statuses[0][1].startswith("rejected:")


async def test_pipeline_still_renders_a_real_docx(monkeypatch):
    rendered = []

    def _render(b):
        rendered.append(b)
        return [PageImage(1, Image.new("RGB", (10, 10), "white"))], ""

    async def _noop(*a, **k):
        return None

    async def _no_questions(self, images, page_infos=None):
        return []

    blob = _real_docx()
    monkeypatch.setattr(pipeline, "prune_debug_crops", lambda *a, **k: 0)
    monkeypatch.setattr(pipeline, "_set_project_status", _noop)
    monkeypatch.setattr(pipeline, "docx_to_images", _render)
    monkeypatch.setattr(pipeline.AIAnalyzer, "extract_all_questions", _no_questions)

    res = await pipeline.process_file(blob, "docx", "pid")

    assert rendered == [blob]
    assert res.status == "no_questions"


# ── Unsafe archive shapes ────────────────────────────────────────────────────

def _symlink_entry():
    info = zipfile.ZipInfo("word/media/link.png")
    info.external_attr = (stat.S_IFLNK | 0o777) << 16
    return info


def _bzip2_entry():
    info = zipfile.ZipInfo("word/media/b.bin")
    info.compress_type = zipfile.ZIP_BZIP2
    return info


def _header_offsets(blob: bytearray, name: bytes) -> tuple[int, int]:
    """(local header, central directory record) offsets of the entry `name`."""
    local = blob.rfind(b"PK\x03\x04", 0, blob.find(name))
    central = blob.rfind(b"PK\x01\x02", 0, blob.rfind(name))
    return local, central


_EXTERNAL_IMAGE_RELS = (
    b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    b'<Relationship Id="rId99" TargetMode="External" Target="http://198.51.100.7/x.png"'
    b' Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"/>'
    b"</Relationships>"
)


@pytest.mark.parametrize("extra,reason", [
    ([("../evil.xml", b"x")], "unsafe archive path"),
    ([("/abs.xml", b"x")], "unsafe archive path"),
    ([("C:evil.xml", b"x")], "unsafe archive path"),
    ([("WORD/document.xml", b"x")], "duplicate archive entry"),
    ([_symlink_entry()], "symbolic link"),
    ([_bzip2_entry()], "unsupported archive compression"),
    ([("word/evil.xml", b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><x>&a;</x>')],
     "DTD"),
    ([("word/_rels/extra.xml.rels", _EXTERNAL_IMAGE_RELS)], "external resource"),
], ids=["dotdot", "absolute", "drive", "duplicate", "symlink", "bzip2", "dtd",
        "external-image"])
def test_unsafe_archives_are_rejected(extra, reason):
    with pytest.raises(InvalidDocx, match=reason):
        validate_docx(_rebuild(_real_docx(), extra=extra))


def test_backslash_entry_name_is_rejected():
    # zipfile normalises "\" to "/" when WRITING on Windows, so plant the
    # backslash by patching the stored name bytes (same length).
    blob = _rebuild(_real_docx(), extra=[("word|evil.xml", b"x")])
    with pytest.raises(InvalidDocx, match="unsafe archive path"):
        validate_docx(blob.replace(b"word|evil.xml", b"word\\evil.xml"))


def test_encrypted_entry_is_rejected():
    # zipfile resets flag bits when writing, so set the encryption bit in
    # both headers afterwards (local: offset 6; central: offset 8).
    blob = bytearray(_rebuild(_real_docx(), extra=[("word/media/e.bin", b"x" * 100)]))
    local, central = _header_offsets(blob, b"word/media/e.bin")
    blob[local + 6] |= 0x1
    blob[central + 8] |= 0x1
    with pytest.raises(InvalidDocx, match="encrypted"):
        validate_docx(bytes(blob))


def test_dtd_split_across_read_chunks_is_found(monkeypatch):
    monkeypatch.setattr(docx_limits, "_CHUNK", 7)
    blob = _rebuild(_real_docx(), extra=[("word/evil.xml", b"<?xml?>  <!DOCTYPE x><x/>")])
    with pytest.raises(InvalidDocx, match="DTD"):
        validate_docx(blob)


def test_too_many_entries(monkeypatch):
    monkeypatch.setattr(docx_limits, "MAX_ENTRIES", 10)
    blob = _rebuild(_real_docx(), extra=[(f"word/media/{i}.bin", b"x") for i in range(10)])
    with pytest.raises(InvalidDocx, match="too many"):
        validate_docx(blob)


def test_total_expanded_size_is_bounded(monkeypatch):
    monkeypatch.setattr(docx_limits, "MAX_TOTAL_BYTES", 64 * 1024)
    blob = _rebuild(_real_docx(), extra=[
        (f"word/media/{i}.bin", bytes(range(256)) * 128) for i in range(3)
    ])
    with pytest.raises(InvalidDocx, match="expands too much"):
        validate_docx(blob)


def test_oversized_upload(monkeypatch):
    monkeypatch.setattr(docx_limits, "MAX_COMPRESSED_BYTES", 1000)
    with pytest.raises(InvalidDocx, match="file too large"):
        validate_docx(_real_docx())


def test_header_that_understates_the_size_is_not_trusted():
    blob = bytearray(_rebuild(_real_docx(), extra=[("word/media/big.bin", b"\n" * 100_000)]))
    # Rewrite the declared uncompressed size to 10 bytes in BOTH headers
    # (local header: offset 22; central directory record: offset 24).
    local, central = _header_offsets(blob, b"word/media/big.bin")
    struct.pack_into("<I", blob, local + 22, 10)
    struct.pack_into("<I", blob, central + 24, 10)
    with pytest.raises(InvalidDocx):
        validate_docx(bytes(blob))


def test_not_a_zip():
    with pytest.raises(InvalidDocx, match="malformed"):
        validate_docx(b"%PDF-1.4 not a docx")
