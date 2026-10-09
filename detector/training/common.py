"""Shared paths, file readers and JSONL helpers for the A2 training pipeline.

Pipeline (each step reads the previous step's output under data/work/):

    1. prepare_human   data/human/*  + meta.csv  -> work/human_chunks.jsonl
    2. generate_ai     human_chunks (topic/genre/lang/length ONLY)
                                                 -> work/ai_generated.jsonl
       generate_ai manual   data/ai_manual/*     -> work/ai_manual.jsonl
    3. split           all of the above          -> work/split/{train,val,test}.jsonl
    4. augment_ocr     split/train + split/test  -> train_aug.jsonl, test_ocr.jsonl
    5. (Colab)         train_aug + val           -> fine-tuned xlm-roberta -> ONNX
    6. evaluate        model + val/test          -> threshold, metrics, CSV
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path
from typing import Iterable, Iterator

DETECTOR_DIR = Path(__file__).resolve().parents[1]
DATA_DIR = Path(os.environ.get("DETECTOR_DATA_DIR", DETECTOR_DIR / "data"))


def human_dir(data: Path = None) -> Path:
    return (data or DATA_DIR) / "human"


def ai_manual_dir(data: Path = None) -> Path:
    return (data or DATA_DIR) / "ai_manual"


def work_dir(data: Path = None) -> Path:
    d = (data or DATA_DIR) / "work"
    d.mkdir(parents=True, exist_ok=True)
    return d


SUPPORTED_SUFFIXES = {".txt", ".docx", ".pdf"}


def read_document(path: Path) -> str:
    """Plain text of a .txt / .docx / .pdf file. Paragraphs are separated by a
    blank line so later steps can chunk on paragraph boundaries."""
    suffix = path.suffix.lower()
    if suffix == ".txt":
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251", "latin-1"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    if suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        paras = [p.text.strip() for p in doc.paragraphs]
        for table in doc.tables:  # coursework often keeps text in tables
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paras.append(cell.text.strip())
        return "\n\n".join(p for p in paras if p)
    if suffix == ".pdf":
        import fitz

        with fitz.open(str(path)) as doc:
            pages = []
            for page in doc:
                blocks = page.get_text("blocks")
                pages.append("\n\n".join(
                    b[4].strip().replace("\n", " ") for b in blocks
                    if len(b) > 4 and str(b[4]).strip()
                ))
        return "\n\n".join(pages)
    raise ValueError(f"unsupported file type: {path.name}")


def read_meta_csv(path: Path) -> dict[str, dict[str, str]]:
    """meta.csv keyed by file name. Column names are case-insensitive;
    `;` and `,` delimiters both work (Excel in uz/ru locales writes `;`)."""
    if not path.exists():
        return {}
    text = path.read_text(encoding="utf-8-sig")
    dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=";,")
    rows = csv.DictReader(text.splitlines(), dialect=dialect)
    out = {}
    for row in rows:
        row = {(k or "").strip().lower(): (v or "").strip() for k, v in row.items()}
        name = row.get("fayl") or row.get("file")
        if name:
            out[name] = row
    return out


def write_jsonl(path: Path, rows: Iterable[dict]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
            n += 1
    return n


def append_jsonl(path: Path, row: dict) -> None:
    """One row, flushed and fsynced: a crash or Ctrl-C loses at most the row
    being written, which is what makes generate_ai resumable."""
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(row, ensure_ascii=False) + "\n")
        f.flush()
        os.fsync(f.fileno())


def read_jsonl(path: Path) -> Iterator[dict]:
    if not path.exists():
        return
    with path.open(encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                yield json.loads(line)
            except json.JSONDecodeError:
                # A half-written last line after a crash: skip it, the row
                # will simply be regenerated.
                continue


def _read_env_var(path: Path, name: str) -> str:
    if not path or not Path(path).is_file():
        return ""
    for line in Path(path).read_text(encoding="utf-8", errors="replace").splitlines():
        if line.strip().startswith(f"{name}="):
            return line.split("=", 1)[1].strip().strip('"').strip("'")
    return ""


def resolve_gen_key(env_file: Path | None = None) -> str:
    """DETECTOR_GEN_GEMINI_KEY from the environment, else from `env_file`
    (only read, never written). Exits if missing or equal to the bot's own
    GEMINI_API_KEY: the training pipeline must never spend the bot's quota."""
    import sys

    key = os.environ.get("DETECTOR_GEN_GEMINI_KEY", "").strip()
    bot_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if env_file:
        key = key or _read_env_var(env_file, "DETECTOR_GEN_GEMINI_KEY")
        bot_key = bot_key or _read_env_var(env_file, "GEMINI_API_KEY")
    if not key:
        sys.exit("DETECTOR_GEN_GEMINI_KEY is not set (environment or --env-file). "
                 "Use a separate FREE key, never the bot's GEMINI_API_KEY.")
    if bot_key and key == bot_key:
        sys.exit("DETECTOR_GEN_GEMINI_KEY equals the bot's GEMINI_API_KEY; use a separate free key.")
    return key


def stable_hash(*parts: str) -> int:
    h = hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()
    return int(h[:12], 16)
